from pathlib import Path
from math import sin, pi

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUTPUT = ROOT / "Apresentacao_Comercial_POC_Radio_Sagres.pptx"

W = 13.333
H = 7.5

RED = RGBColor(152, 59, 46)
YELLOW = RGBColor(252, 175, 38)
ORANGE = RGBColor(245, 130, 41)
CLAY = RGBColor(214, 116, 48)
EMBER = RGBColor(221, 89, 45)
CREAM = RGBColor(255, 249, 241)
INK = RGBColor(54, 35, 31)
MUTED = RGBColor(112, 91, 86)
WHITE = RGBColor(255, 255, 255)
LIGHT = RGBColor(247, 237, 232)
FONT = "Calibri"


def add_shape(slide, kind, x, y, w, h, fill, line=None, radius=True):
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line if line else fill
    return shape


def add_text(slide, text, x, y, w, h, size=20, color=INK, bold=False,
             align=PP_ALIGN.LEFT, font=FONT, valign=MSO_ANCHOR.TOP,
             margin=0.04, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def add_rich_text(slide, chunks, x, y, w, h, size=20, color=INK,
                  align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    for text, bold, chunk_color in chunks:
        run = p.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = chunk_color or color
    return box


def add_picture_crop(slide, path, x, y, w, h):
    from PIL import Image

    with Image.open(path) as image:
        image_ratio = image.width / image.height
    frame_ratio = w / h
    picture = slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                       width=Inches(w), height=Inches(h))
    if image_ratio > frame_ratio:
        shown = frame_ratio / image_ratio
        crop = (1 - shown) / 2
        picture.crop_left = crop
        picture.crop_right = crop
    elif image_ratio < frame_ratio:
        shown = image_ratio / frame_ratio
        crop = (1 - shown) / 2
        picture.crop_top = crop
        picture.crop_bottom = crop
    return picture


def add_background(slide, color=CREAM):
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, W, H, color)


def add_brand_mark(slide, x=10.95, y=0.27, w=1.95, h=0.85):
    add_picture_crop(slide, ASSETS / "sagres-am730-logo.jpg", x, y, w, h)


def add_footer(slide, number, dark=False):
    color = WHITE if dark else MUTED
    add_text(slide, "POC APP SAGRES • APRESENTAÇÃO COMERCIAL", 0.55, 7.12,
             5.8, 0.2, 8, color, bold=True)
    add_text(slide, f"{number:02d}", 12.15, 7.06, 0.55, 0.24, 9, color,
             bold=True, align=PP_ALIGN.RIGHT)


def add_title(slide, kicker, title, subtitle=None, dark=False, number=1):
    title_color = WHITE if dark else INK
    subtitle_color = RGBColor(255, 222, 211) if dark else MUTED
    add_text(slide, kicker.upper(), 0.58, 0.42, 6.6, 0.28, 10,
             YELLOW if dark else RED, bold=True)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0.58, 0.78, 0.58, 0.045,
              YELLOW)
    add_text(slide, title, 0.58, 0.98, 8.7, 0.82, 28, title_color, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.58, 1.78, 8.7, 0.54, 14,
                 subtitle_color)
    add_brand_mark(slide)
    add_footer(slide, number, dark)


def add_card(slide, x, y, w, h, title, body, accent=YELLOW,
             fill=WHITE, icon=None, dark=False):
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, fill)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
              x + 0.18, y + 0.2, 0.42, 0.42, accent)
    if icon:
        add_text(slide, icon, x + 0.18, y + 0.19, 0.42, 0.42, 17,
                 RED if accent == YELLOW else WHITE, bold=True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text_color = WHITE if dark else INK
    body_color = RGBColor(255, 226, 216) if dark else MUTED
    add_text(slide, title, x + 0.18, y + 0.76, w - 0.36, 0.42, 16,
             text_color, bold=True)
    add_text(slide, body, x + 0.18, y + 1.18, w - 0.36, h - 1.32, 11,
             body_color)


def add_phone(slide, image, x, y, h, label=None, accent=YELLOW):
    ratio = 1080 / 2400
    w = h * ratio
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
              x - 0.09, y - 0.09, w + 0.18, h + 0.18, INK)
    add_picture_crop(slide, image, x, y, w, h)
    if label:
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                  x + 0.16, y + h - 0.48, w - 0.32, 0.32, accent)
        add_text(slide, label.upper(), x + 0.16, y + h - 0.46,
                 w - 0.32, 0.25, 9, RED, bold=True, align=PP_ALIGN.CENTER)
    return w


def add_wave(slide, y, color=YELLOW, width=1.8, amplitude=0.1):
    points = []
    for i in range(25):
        x = -0.2 + i * (W + 0.4) / 24
        yy = y + sin(i / 24 * 4 * pi) * amplitude
        points.append((x, yy))
    for a, b in zip(points, points[1:]):
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                          Inches(a[0]), Inches(a[1]),
                                          Inches(b[0]), Inches(b[1]))
        line.line.color.rgb = color
        line.line.width = Pt(width)


def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, RED)
    add_wave(slide, 6.45, YELLOW, 2.5, 0.11)
    add_wave(slide, 6.67, ORANGE, 2.5, 0.08)
    add_picture_crop(slide, ASSETS / "sagres-am730-logo.jpg", 0.6, 0.45, 2.65, 1.28)
    add_text(slide, "SAGRES NO BOLSO.", 0.65, 2.12, 7.2, 0.78, 33, WHITE, bold=True)
    add_text(slide, "Ao vivo, em tom maior.", 0.65, 2.9, 6.3, 0.5, 22, YELLOW, bold=True)
    add_text(slide,
             "Uma proposta de canal digital próprio para aproximar audiência, conteúdo e marca.",
             0.65, 3.62, 5.85, 1.0, 17, RGBColor(255, 226, 216))
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
              8.3, 0.58, 4.18, 6.35, INK)
    add_phone(slide, ASSETS / "app-playing.png", 9.12, 0.82, 5.88)
    add_text(slide, "POC FLUTTER • JULHO 2026", 0.65, 6.88, 4.3, 0.23, 9,
             WHITE, bold=True)


def slide_opportunity(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title(slide, "Oportunidade", "O áudio já é digital.\nA relação também pode ser.",
              "Hoje, o stream entrega conteúdo. O aplicativo pode entregar experiência, vínculo e recorrência.",
              number=2)
    items = [
        ("01", "Acesso direto", "Um ícone da Sagres na tela do ouvinte, sem depender da jornada de terceiros."),
        ("02", "Presença contínua", "Áudio em segundo plano mantém a marca presente no cotidiano."),
        ("03", "Base para crescer", "Programação, notícias, podcasts e relacionamento podem evoluir no mesmo canal."),
    ]
    for i, (num, title, body) in enumerate(items):
        x = 0.62 + i * 4.15
        add_card(slide, x, 3.15, 3.72, 2.55, title, body,
                 [YELLOW, ORANGE, EMBER][i], icon=num)
    add_text(slide, "CANAL PRÓPRIO  →  EXPERIÊNCIA  →  RELACIONAMENTO", 2.2, 6.23,
             8.9, 0.4, 15, RED, bold=True, align=PP_ALIGN.CENTER)


def slide_proposal(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, RED)
    add_title(slide, "Proposta", "Uma porta de entrada simples\npara todo o universo Sagres.",
              "Começar pelo hábito mais forte — ouvir ao vivo — e evoluir conforme a estratégia de conteúdo.",
              dark=True, number=3)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.72, 3.12, 11.9, 2.5,
              RGBColor(123, 42, 33))
    labels = [
        ("RÁDIO AO VIVO", "já validado"),
        ("PROGRAMAÇÃO", "experiência pronta"),
        ("NOTÍCIAS", "experiência pronta"),
        ("OUVINTE", "visão futura"),
    ]
    for i, (title, status) in enumerate(labels):
        x = 1.05 + i * 2.92
        color = YELLOW if i == 0 else ORANGE if i < 3 else EMBER
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, 3.55, 2.35, 1.3, color)
        add_text(slide, title, x + 0.1, 3.74, 2.15, 0.36, 13, RED, bold=True,
                 align=PP_ALIGN.CENTER)
        add_text(slide, status, x + 0.1, 4.18, 2.15, 0.24, 10, RED,
                 align=PP_ALIGN.CENTER)
        if i < 3:
            add_text(slide, "→", x + 2.43, 3.91, 0.42, 0.4, 22, WHITE, bold=True,
                     align=PP_ALIGN.CENTER)
    add_rich_text(slide,
                  [("Começar pequeno. ", True, YELLOW),
                   ("Validar rápido. ", True, WHITE),
                   ("Evoluir com evidências.", True, YELLOW)],
                  2.0, 6.16, 9.3, 0.5, 20, WHITE, PP_ALIGN.CENTER)


def slide_validated(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title(slide, "Prova de conceito", "O que já foi validado",
              "Não é apenas uma tela: o fluxo principal funciona de ponta a ponta.",
              number=4)
    cards = [
        ("▶", "Stream real", "Reprodução do áudio oficial Sagres em MP3."),
        ("◼", "Segundo plano", "Controles de mídia e continuidade no sistema."),
        ("≈", "Estado vivo", "Ondas, equalizador e glow seguem o áudio real."),
        ("↻", "Resiliência", "Loading, falha de rede e tentativa novamente."),
        ("⌂", "Quatro áreas", "Início, Programação, Notícias e Você."),
        ("✓", "Qualidade", "Arquitetura testável, suíte local e fluxo E2E."),
    ]
    for i, (icon, title, body) in enumerate(cards):
        row, col = divmod(i, 3)
        add_card(slide, 0.62 + col * 4.15, 2.65 + row * 1.82, 3.72, 1.55,
                 title, body, [YELLOW, ORANGE, EMBER][col], icon=icon)


def slide_experience(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, RED)
    add_title(slide, "Experiência", "O estado do áudio\ntambém é visual.",
              "A interface comunica com clareza quando está tocando e quando está pausada.",
              dark=True, number=5)
    add_phone(slide, ASSETS / "app-playing.png", 7.22, 1.2, 5.75, "tocando")
    add_phone(slide, ASSETS / "app-paused.png", 10.15, 1.2, 5.75, "pausada", ORANGE)
    bullets = [
        "Ondas e equalizador animados em reprodução",
        "Botão amarelo pulsa como feedback de atividade",
        "Todas as animações param ao pausar",
        "Identidade inspirada no Cerrado e na navegação",
    ]
    for i, text in enumerate(bullets):
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 0.75, 3.08 + i * 0.72,
                  0.25, 0.25, YELLOW)
        add_text(slide, text, 1.15, 2.96 + i * 0.72, 5.45, 0.47, 15, WHITE,
                 bold=i < 3)


def slide_journey(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title(slide, "Jornada", "Três toques. Uma relação contínua.",
              "A proposta reduz atrito e torna a Sagres acessível em qualquer momento.",
              number=6)
    steps = [
        ("1", "ABRIR", "A marca já está na tela inicial do celular."),
        ("2", "OUVIR", "Um toque inicia a transmissão oficial."),
        ("3", "CONTINUAR", "O áudio segue no bolso, na tela bloqueada e no fone."),
    ]
    for i, (num, title, body) in enumerate(steps):
        x = 0.72 + i * 4.18
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, 2.8, 3.55, 2.65,
                  WHITE)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 0.28, 3.12, 0.76, 0.76,
                  [YELLOW, ORANGE, EMBER][i])
        add_text(slide, num, x + 0.28, 3.18, 0.76, 0.52, 23, RED, bold=True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_text(slide, title, x + 0.3, 4.08, 2.9, 0.35, 16, RED, bold=True)
        add_text(slide, body, x + 0.3, 4.52, 2.9, 0.7, 12, MUTED)
        if i < 2:
            add_text(slide, "→", x + 3.62, 3.78, 0.45, 0.45, 24, RED, bold=True,
                     align=PP_ALIGN.CENTER)


def slide_listener_value(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, RED)
    add_title(slide, "Valor para o ouvinte", "Mais fácil ouvir.\nMais fácil permanecer.",
              "A tecnologia fica em segundo plano para o conteúdo ocupar o centro da experiência.",
              dark=True, number=7)
    cards = [
        ("◎", "Conveniência", "Acesso direto ao vivo, sem procurar a página certa."),
        ("♫", "Continuidade", "Áudio em segundo plano durante outras atividades."),
        ("▣", "Contexto", "Programa atual e próximos conteúdos no mesmo lugar."),
        ("♡", "Proximidade", "Uma base futura para favoritos e relacionamento."),
    ]
    for i, (icon, title, body) in enumerate(cards):
        add_card(slide, 0.72 + i * 3.1, 3.0, 2.68, 2.55, title, body,
                 [YELLOW, ORANGE, CLAY, EMBER][i], icon=icon,
                 fill=RGBColor(123, 42, 33), dark=True)


def slide_business_value(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title(slide, "Valor para a Sagres", "Um ativo digital próprio",
              "A POC abre espaço para uma estratégia de produto, conteúdo e relacionamento.",
              number=8)
    items = [
        ("MARCA", "Presença recorrente no dispositivo do ouvinte."),
        ("CONTEÚDO", "Distribuição integrada de rádio, notícias e podcasts."),
        ("RELACIONAMENTO", "Canal direto para preferências e comunicação futura."),
        ("APRENDIZADO", "Base para entender jornadas e melhorar a oferta."),
    ]
    for i, (title, body) in enumerate(items):
        row, col = divmod(i, 2)
        x, y = 0.72 + col * 6.12, 2.72 + row * 1.72
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, 5.55, 1.38,
                  WHITE)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, 0.12, 1.38,
                  [YELLOW, ORANGE, CLAY, EMBER][i])
        add_text(slide, title, x + 0.38, y + 0.25, 2.2, 0.3, 13, RED, bold=True)
        add_text(slide, body, x + 2.2, y + 0.2, 2.95, 0.72, 12, MUTED)
    add_text(slide, "A POC valida a fundação. O retorno comercial depende da estratégia de produto e operação.",
             1.2, 6.3, 10.9, 0.42, 12, RED, bold=True, align=PP_ALIGN.CENTER)


def slide_commercial(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, RED)
    add_title(slide, "Potencial comercial", "Novas superfícies,\nsem interromper a experiência.",
              "Possibilidades a validar com programação, comercial e audiência.",
              dark=True, number=9)
    options = [
        ("Programa apresentado por", "Assinatura contextual na tela do programa."),
        ("Conteúdo patrocinado", "Cards editoriais claramente identificados."),
        ("Podcast de marca", "Séries especiais disponíveis sob demanda."),
        ("Campanha local", "Ativações conectadas a Goiás e à comunidade."),
    ]
    for i, (title, body) in enumerate(options):
        add_card(slide, 0.72 + i * 3.1, 3.0, 2.68, 2.55, title, body,
                 [YELLOW, ORANGE, CLAY, EMBER][i], icon=str(i + 1),
                 fill=RGBColor(123, 42, 33), dark=True)
    add_text(slide, "HIPÓTESES COMERCIAIS • NÃO INCLUÍDAS NA POC ATUAL", 2.0, 6.13,
             9.3, 0.33, 11, YELLOW, bold=True, align=PP_ALIGN.CENTER)


def slide_technology(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title(slide, "Tecnologia", "Uma fundação preparada para evoluir",
              "Arquitetura hexagonal separa experiência, regras e integrações.",
              number=10)
    layers = [
        ("EXPERIÊNCIA", "Flutter • telas • animações", YELLOW),
        ("APLICAÇÃO", "estado • comandos • tratamento de falhas", ORANGE),
        ("DOMÍNIO", "estação • player • programação", CLAY),
        ("ADAPTADORES", "áudio nativo • dados em memória • APIs futuras", EMBER),
    ]
    for i, (title, body, color) in enumerate(layers):
        y = 2.6 + i * 0.83
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                  0.75 + i * 0.3, y, 7.05 - i * 0.6, 0.62, color)
        add_text(slide, title, 1.0 + i * 0.3, y + 0.1, 1.55, 0.25, 11, RED, bold=True)
        add_text(slide, body, 2.55 + i * 0.2, y + 0.1, 4.6 - i * 0.55, 0.25,
                 11, RED)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 8.65, 2.55, 3.75, 3.55,
              RED)
    points = [
        "Trocar a fonte de dados sem refazer as telas",
        "Testar regras sem depender da internet",
        "Adicionar backend quando houver necessidade real",
        "Manter Android e iOS na mesma base",
    ]
    add_text(slide, "POR QUE IMPORTA", 9.0, 2.9, 3.0, 0.34, 14, YELLOW, bold=True)
    for i, point in enumerate(points):
        add_text(slide, "✓", 9.0, 3.48 + i * 0.56, 0.3, 0.3, 14, YELLOW, bold=True)
        add_text(slide, point, 9.38, 3.42 + i * 0.56, 2.55, 0.46, 11, WHITE)


def slide_quality(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, RED)
    add_title(slide, "Confiança", "Evidências antes da próxima etapa",
              "A POC foi construída para demonstrar e também para ser verificada.",
              dark=True, number=11)
    metrics = [
        ("12", "testes automatizados", "regras, estado visual e navegação"),
        ("E2E", "fluxo em Android", "play, pause e quatro áreas"),
        ("REAL", "stream oficial", "MP3 em reprodução no emulador"),
        ("2", "plataformas", "Android e iOS preparados"),
    ]
    for i, (value, label, detail) in enumerate(metrics):
        x = 0.72 + i * 3.1
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, 2.9, 2.68, 2.65,
                  RGBColor(123, 42, 33))
        add_text(slide, value, x + 0.15, 3.18, 2.38, 0.68, 31, YELLOW, bold=True,
                 align=PP_ALIGN.CENTER)
        add_text(slide, label, x + 0.18, 4.03, 2.32, 0.4, 14, WHITE, bold=True,
                 align=PP_ALIGN.CENTER)
        add_text(slide, detail, x + 0.2, 4.55, 2.28, 0.62, 10,
                 RGBColor(255, 222, 211), align=PP_ALIGN.CENTER)


def slide_roadmap(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title(slide, "Roadmap", "Da POC ao produto",
              "Uma evolução em etapas mantém investimento e aprendizado sob controle.",
              number=12)
    phases = [
        ("01", "CONTEÚDO REAL", "Grade oficial\nNotícias e podcasts\nConfiguração remota", YELLOW),
        ("02", "PILOTO CONTROLADO", "Teste com ouvintes\nMétricas essenciais\nAjustes de experiência", ORANGE),
        ("03", "PRODUTO", "Backend e perfil\nNotificações\nOperação e lojas", EMBER),
    ]
    for i, (num, title, body, color) in enumerate(phases):
        x = 0.72 + i * 4.15
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, 2.8, 3.72, 3.1,
                  WHITE)
        add_text(slide, num, x + 0.2, 3.02, 0.75, 0.48, 23, color, bold=True)
        add_text(slide, title, x + 0.2, 3.68, 3.1, 0.38, 15, RED, bold=True)
        add_text(slide, body, x + 0.2, 4.27, 3.1, 1.18, 13, MUTED)
        if i < 2:
            add_text(slide, "→", x + 3.75, 4.02, 0.38, 0.42, 23, RED, bold=True,
                     align=PP_ALIGN.CENTER)


def slide_decision(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, RED)
    add_wave(slide, 6.5, YELLOW, 2.5, 0.12)
    add_wave(slide, 6.72, ORANGE, 2.5, 0.08)
    add_picture_crop(slide, ASSETS / "sagres-am730-logo.jpg", 0.62, 0.46, 2.5, 1.18)
    add_text(slide, "A próxima decisão", 0.68, 2.0, 5.4, 0.5, 17, YELLOW, bold=True)
    add_text(slide, "Transformar a POC\nem um piloto de produto.", 0.68, 2.62, 7.3, 1.45,
             34, WHITE, bold=True)
    add_text(slide,
             "Alinhar conteúdo real, escopo do piloto e critérios de sucesso com Produto, Tecnologia, Programação e Comercial.",
             0.68, 4.38, 6.35, 1.0, 17, RGBColor(255, 226, 216))
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 8.05, 1.58, 4.35, 4.55,
              RGBColor(123, 42, 33))
    add_text(slide, "PRÓXIMOS PASSOS", 8.5, 2.02, 3.35, 0.4, 15, YELLOW, bold=True)
    steps = [
        "1. Validar a proposta com áreas-chave",
        "2. Confirmar fontes de conteúdo",
        "3. Definir público e duração do piloto",
        "4. Priorizar evolução técnica",
        "5. Colocar a experiência em campo",
    ]
    for i, step in enumerate(steps):
        add_text(slide, step, 8.5, 2.72 + i * 0.57, 3.35, 0.38, 12, WHITE,
                 bold=i == 4)
    add_text(slide, "SAGRES • EM TOM MAIOR", 0.68, 6.92, 3.5, 0.25, 10, WHITE,
             bold=True)


def slide_scope(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title(slide, "Apêndice", "O que é POC e o que é produto",
              "Transparência para orientar decisão e investimento.", number=14)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.72, 2.63, 5.72, 3.2,
              WHITE)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 6.86, 2.63, 5.72, 3.2,
              RED)
    add_text(slide, "POC ATUAL", 1.05, 2.95, 4.9, 0.4, 17, RED, bold=True)
    add_text(slide,
             "✓ Rádio ao vivo\n✓ Segundo plano\n✓ Estados e animações\n✓ Navegação demonstrativa\n✓ Testes automatizados",
             1.05, 3.55, 4.75, 1.75, 14, MUTED)
    add_text(slide, "PRODUTO FUTURO", 7.2, 2.95, 4.9, 0.4, 17, YELLOW, bold=True)
    add_text(slide,
             "○ Conteúdo editorial real\n○ Backend e configuração remota\n○ Perfil e preferências\n○ Métricas e notificações\n○ Operação, segurança e lojas",
             7.2, 3.55, 4.75, 1.75, 14, WHITE)
    add_text(slide,
             "Premissa: qualquer projeção comercial depende da definição do produto, audiência e modelo de operação.",
             1.3, 6.35, 10.7, 0.45, 12, RED, bold=True, align=PP_ALIGN.CENTER)


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "Sagres no bolso — POC App Rádio Sagres"
    prs.core_properties.subject = "Apresentação comercial e didática da POC Flutter"
    prs.core_properties.author = "Instituto Sagres"
    prs.core_properties.keywords = "Sagres, rádio, Flutter, POC, aplicativo"

    slide_cover(prs)
    slide_opportunity(prs)
    slide_proposal(prs)
    slide_validated(prs)
    slide_experience(prs)
    slide_journey(prs)
    slide_listener_value(prs)
    slide_business_value(prs)
    slide_commercial(prs)
    slide_technology(prs)
    slide_quality(prs)
    slide_roadmap(prs)
    slide_decision(prs)
    slide_scope(prs)

    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
